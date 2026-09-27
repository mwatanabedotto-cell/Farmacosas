"""Conector de AGEMED (Bolivia): Lista Nacional de Medicamentos Esenciales (LINAME).

AGEMED no publica de forma abierta el registro sanitario (marcas y números de registro): su buscador
exige reCAPTCHA y no se automatiza. Sí publica la LINAME vigente en Excel, que indica qué medicamentos
(DCI, forma y concentración) están en la lista oficial, cuáles son de uso restringido y la clasificación
AWaRe de la OMS de los antibióticos.
"""

import io
import re
from collections.abc import Callable
from urllib.parse import urljoin, urlparse

import httpx
from openpyxl import load_workbook

from app.connectors.base import ClienteHTTP, ErrorConector

URL_WEB = "https://www.agemed.gob.bo/"
URL_PAGINA = "https://www.agemed.gob.bo/#dtu/contenido"
# El sitio carga sus secciones desde esta API; la sección de uso racional enlaza la LINAME.
URL_SECCION = "https://apiwww.agemed.gob.bo/api/web/dtu|contenido"
ENLACE_LINAME = re.compile(
    r"archivo_uso_racional/liname/LINAME_(\d{4})_(\d{4})(?:_(\d{2})-(\d{2})-(\d{2}))?[^\"'\\\s]*\.xlsx", re.I
)
MAX_BYTES = 20 * 1024 * 1024
ATC = re.compile(r"[A-Z]\d{2}[A-Z]{2}\d{2}")


class AgemedError(ErrorConector):
    pass


def elegir_liname(html: str) -> str:
    """URL absoluta de la LINAME más reciente enlazada en la página (por años y fecha de actualización)."""
    candidatos = []
    for m in ENLACE_LINAME.finditer(html.replace("\\/", "/")):
        anio1, anio2, dia, mes, anio = m.groups()
        fecha = f"20{anio}-{mes}-{dia}" if anio else ""
        candidatos.append(((int(anio1), int(anio2), fecha), m.group(0)))
    if not candidatos:
        raise AgemedError("No se encontró el enlace a la LINAME en la web de AGEMED")
    url = urljoin(URL_WEB, max(candidatos)[1])
    if urlparse(url).hostname != urlparse(URL_WEB).hostname:
        raise AgemedError(f"Enlace fuera de agemed.gob.bo: {url}")
    return url


def _texto(valor) -> str:
    """Celda a texto; los porcentajes de Excel llegan como fracción (0.05 -> '5 %')."""
    if valor is None:
        return ""
    if isinstance(valor, float) and 0 < valor < 1:
        return f"{valor * 100:g} %"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).split())


def _codigo(fila) -> str | None:
    letra, grupo, correlativo = (fila + (None,) * 3)[:3]
    if not (isinstance(letra, str) and re.fullmatch(r"[A-Z]", letra.strip())):
        return None
    try:
        return f"{letra.strip()}.{int(grupo):02d}.{int(correlativo):02d}"
    except (TypeError, ValueError):
        return None


def leer_liname(contenido: bytes) -> tuple[str, list[dict]]:
    """Devuelve (nombre de la lista, entradas) a partir del Excel de la LINAME."""
    if not contenido.startswith(b"PK"):
        raise AgemedError("El archivo de la LINAME no es un Excel válido")
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    principal = next((h for h in libro.worksheets if h.title.upper().startswith("LINAME") and "ATQ" not in h.title.upper()), None)
    if principal is None:
        raise AgemedError("No se encontró la hoja principal de la LINAME")

    aware: dict[str, str] = {}
    hoja_aware = next((h for h in libro.worksheets if "AWARE" in h.title.upper()), None)
    for fila in hoja_aware.iter_rows(values_only=True) if hoja_aware else []:
        codigo = _codigo(tuple(fila[1:4]))
        categoria = _texto(fila[8]) if len(fila) > 8 else ""
        if codigo and categoria:
            aware[codigo] = categoria.capitalize()

    entradas = []
    for fila in principal.iter_rows(values_only=True):
        codigo = _codigo(tuple(fila[:3]))
        if not codigo or len(fila) < 7 or not _texto(fila[3]):
            continue
        atc = _texto(fila[6]).upper().replace(" ", "")
        entradas.append({
            "codigo": codigo,
            "medicamento": _texto(fila[3]),
            "forma_farmaceutica": _texto(fila[4]).capitalize(),
            "concentracion": _texto(fila[5]),
            "atc": atc or None,
            "uso_restringido": _texto(fila[7] if len(fila) > 7 else None).upper() == "R",
            "aware": aware.get(codigo),
        })
    if not entradas:
        raise AgemedError("La LINAME no contiene entradas reconocibles")
    return principal.title.strip(), entradas


class AgemedClient:
    def __init__(
        self,
        http: httpx.Client | None = None,
        pausa_segundos: float = 0.2,
        max_reintentos: int = 4,
        dormir: Callable[[float], None] | None = None,
    ):
        self.http = http or httpx.Client(timeout=90, follow_redirects=True)
        kwargs = {"dormir": dormir} if dormir else {}
        self.cliente = ClienteHTTP(self.http, pausa_segundos, max_reintentos, error=AgemedError, **kwargs)

    def liname(self) -> tuple[str, str, list[dict]]:
        """Devuelve (url del Excel, nombre de la lista, entradas)."""
        pagina = self.cliente.get(URL_SECCION)
        if pagina is None:
            raise AgemedError("La sección de uso racional de AGEMED no está disponible")
        url = elegir_liname(pagina.text)
        self.cliente.pausa()
        archivo = self.cliente.get(url)
        if archivo is None:
            raise AgemedError(f"No se pudo descargar {url}")
        if len(archivo.content) > MAX_BYTES:
            raise AgemedError("El archivo de la LINAME supera el tamaño máximo esperado")
        nombre, entradas = leer_liname(archivo.content)
        return url, nombre, entradas
